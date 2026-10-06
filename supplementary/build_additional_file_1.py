"""Build the editable supplementary note from checked-in audit summaries."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SUMMARY = json.loads((ROOT / "data/pubchem/seed123_pictograms_rare2x/summary.json").read_text(encoding="utf-8"))
NITE = json.loads((ROOT / "external_validation/nite_japan_ghs/nite_external_summary.json").read_text(encoding="utf-8"))
SOURCE = json.loads((ROOT / "external_validation/nite_japan_ghs/nite_external_source_audit.json").read_text(encoding="utf-8"))
with (HERE / "scaffold_paired_bootstrap_delta.csv").open(encoding="utf-8", newline="") as handle:
    BOOTSTRAP = list(csv.DictReader(handle))

LABELS = [
    "flammable", "oxidizing", "gas_under_pressure", "corrosive",
    "acute_toxicity", "irritant_harmful", "cmr", "stot", "environmental_hazard",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, "E9EEF5")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()


def caption(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = "Caption"
    paragraph.add_run(title).bold = True


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Arial"
    doc.styles["Title"].font.size = Pt(16)
    doc.styles["Heading 1"].font.size = Pt(12)

    doc.add_heading("Supplementary methods and audit tables", 0)
    doc.add_paragraph(
        "ChemHazard-LLM: fixed seed123 PubChem split and the InChIKey-disjoint NITE Japan-GHS external set. "
        "This file contains auditable dataset and statistical details; complete model-wise per-label performance "
        "will be supplied in a separate additional file after the original server evaluation outputs are recovered."
    )

    doc.add_heading("Fixed split and training exposure", 1)
    doc.add_paragraph(
        "The main PubChem experiment used a custom multilabel-stratified split with seed 123. The unaugmented "
        "train, validation and test partitions contained 9,244, 1,156 and 1,156 unique compounds, respectively. "
        "Rare-label oversampling was applied only to the training partition. A training record positive for any "
        "preselected target label was duplicated once; labels on a duplicated record were not changed. The "
        "expanded training file therefore contains 12,990 rows, while validation and test retain their original sizes."
    )
    caption(doc, "Supplementary Table S1. Positive-label counts before and after training oversampling")
    rows = []
    for label in LABELS:
        rows.append([
            label,
            str(SUMMARY["label_counts_train_before_oversampling"][label]),
            str(SUMMARY["label_counts_train"][label]),
            str(SUMMARY["label_counts_val"][label]),
            str(SUMMARY["label_counts_test"][label]),
        ])
    add_table(doc, ["Hazard label", "Train original", "Train exposure", "Validation", "Test"], rows)
    doc.add_paragraph(
        "The prespecified oversampling targets were oxidizing, gas_under_pressure, cmr, stot and "
        "environmental_hazard. Counts in the exposure column include repeated rows; these are not new compounds. "
        "Non-target labels may also gain exposure when they co-occur on a duplicated record. "
        "Source: data/pubchem/seed123_pictograms_rare2x/summary.json."
    )

    doc.add_heading("NITE source and label coverage", 1)
    doc.add_paragraph(
        "The external source was the official NITE Japan-GHS classification spreadsheet listed on the NITE "
        f"download page, updated {SOURCE['nite_latest_page_update_date']}. The downloaded source SHA-256 was "
        f"{SOURCE['nite_local_sha256']}. CAS identifiers were resolved to PubChem structures and collapsed only "
        "when all candidate records shared one full InChIKey. Records overlapping the 11,556-compound seed123 "
        "experiment universe by full InChIKey were removed. The resulting primary set contains 1,053 compounds "
        "and 4,630 scoreable label cells. A stricter connectivity-block-disjoint subset contains 1,016 compounds."
    )
    caption(doc, "Supplementary Table S2. NITE reference-label coverage in the external test set")
    rows = []
    for label in LABELS:
        values = NITE["label_summary_exact_disjoint"][label]
        rows.append([label, str(values["valid"]), str(values["positive"]), str(values["negative"])])
    add_table(doc, ["Hazard label", "Scoreable", "Positive", "Negative"], rows)
    doc.add_paragraph(
        "Only cells with an explicit reference classification were scored. A dash or 'Classification not "
        "possible' remained unknown rather than negative. The primary set has only seven nine-label complete "
        "cases. Acute toxicity and irritant/harmful have no verified negative cells under this mapping, so "
        "label-specific specificity cannot be estimated for those categories. "
        "Source: external_validation/nite_japan_ghs/nite_external_summary.json."
    )

    doc.add_heading("Scaffold split paired comparison", 1)
    doc.add_paragraph(
        "The paired bootstrap compared the Mistral LoRA single model with Morgan fingerprints plus "
        "physicochemical descriptors in a random forest on the same scaffold test set. The reported difference "
        "is Mistral minus the random-forest baseline. A negative difference favors Mistral for Hamming loss "
        "because lower loss is better."
    )
    caption(doc, "Supplementary Table S3. Paired bootstrap differences on the scaffold test set")
    metric_names = {"subset_accuracy": "Subset accuracy", "micro_f1": "Micro-F1", "macro_f1": "Macro-F1", "hamming_loss": "Hamming loss"}
    rows = []
    for row in BOOTSTRAP:
        rows.append([
            metric_names[row["metric"]],
            f"{float(row['delta_mistral_minus_baseline']):+.4f}",
            f"[{float(row['ci95_low']):+.4f}, {float(row['ci95_high']):+.4f}]",
        ])
    add_table(doc, ["Metric", "Difference", "95% confidence interval"], rows)
    doc.add_paragraph(
        "Source: supplementary/scaffold_paired_bootstrap_delta.csv, copied from the project's "
        "paper_assets/scaffold_paired_bootstrap_delta.csv. This comparison concerns the single model; "
        "the three-model majority vote did not exceed it on the scaffold split."
    )

    doc.add_heading("Decision rule and scope", 1)
    doc.add_paragraph(
        "The final three-member vote used training seeds 123, 777 and 2025 on the fixed seed123 split. "
        "Validation-set selection yielded a vote threshold of one for oxidizing, gas_under_pressure and cmr, "
        "and two for the other six labels. These thresholds were fixed before evaluation on the PubChem test "
        "set and NITE external set. NITE results are mask-aware and are not directly comparable in absolute "
        "value with fully observed PubChem results. The historical ICSC experiment is excluded from this "
        "supplement because it was stopped."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("ChemHazard-LLM supplementary material | working version")

    output = HERE / "Additional_file_1_methods_and_audits.docx"
    doc.save(output)
    return output


if __name__ == "__main__":
    print(build())
